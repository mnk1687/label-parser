import tempfile,unittest
from pathlib import Path
import fitz
import test_batch
from label_parser.workflow import discover,recognize
from label_parser.prodsheet import ProductionRow
from label_parser.composition import make_compositions,compose,sorted_rows,safe_name

class CompositionTest(unittest.TestCase):
    def test_same_barcode_combines_original_objects_and_preserves_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);f=test_batch.BatchTest()
            f.sheet(root/'ProdSheet.pdf');f.label(root/'FB_MAIN.pdf',test_batch.A,'5208133300829');f.label(root/'FB_CARE.pdf','','5208133300829')
            
            for kind in ['COMPO1','COMPO2','COMPO3']:f.label(root/f'FB_{kind}.pdf','','5208133300829')
            jobs,issues=discover([str(root)],lambda _:None);self.assertFalse(issues)
            job=jobs[0];recognize(job,lambda _:None)
            rows=[ProductionRow('CARELABEL_2-','26',test_batch.A,test_batch.A+'-BLACK-26','5208133300829',54)]
            item=make_compositions(job,rows)[0];self.assertFalse(item.missing);self.assertEqual(len(item.matches),5)
            before={p:p.read_bytes() for p in root.glob('*.pdf')}
            with compose(item) as out:
                self.assertEqual(len(out),1);self.assertEqual(tuple(out[0].rect),(0,0,612,792))
                self.assertTrue(out[0].get_drawings());self.assertFalse(out[0].get_images())
                self.assertEqual(out.metadata['title'],rows[0].txt2)
            for p,data in before.items():self.assertEqual(p.read_bytes(),data)
            missing=make_compositions(job,[ProductionRow('', '27',test_batch.A,'same-model-other-size','5208133300836',1)])[0]
            self.assertTrue(missing.missing)
            with self.assertRaises(ValueError):compose(missing)
    def test_txt1_color_and_size_order(self):
        rows=[ProductionRow('',s,m,m+'-'+color+'-'+s,'5208133300829',1) for m,color,s in [('FBL015-162-03','BLACK','26'),('FBL015-161-02','MD BLUE','25'),('FBL015-161-02','LT BLUE','31'),('FBL015-161-02','LT BLUE','25')]]
        sorted_=sorted_rows(rows)
        self.assertEqual([r.txt2 for r in sorted_],['FBL015-161-02-LT BLUE-25','FBL015-161-02-LT BLUE-31','FBL015-161-02-MD BLUE-25','FBL015-162-03-BLACK-26'])
        self.assertEqual(safe_name('A/B:C'),'A_B_C')

class SummaryTest(unittest.TestCase):
    def test_sorted_vector_pages_and_no_overwrite(self):
        from label_parser.composition import export_summary
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);entries=[];before={};pixels={}
            for index,(color,size) in enumerate([('LT BLUE','40'),('DK BLUE','40'),('DK BLUE','28'),('LT BLUE','28')]):
                row=ProductionRow('',size,'FBM015-099-03',f'FBM015-099-03-{color}-{size}',str(index),1)
                path=root/f'{index}.pdf'
                with fitz.open() as doc:
                    page=doc.new_page(width=612,height=792);page.insert_text((40,40),row.txt2)
                    page.draw_rect(fitz.Rect(40,60,100+index,120),color=(0,1,1))
                    doc.save(path)
                    pixels[row.txt2]=page.get_pixmap().samples
                entries.append((row,path));before[path]=path.read_bytes()
            output=export_summary(entries,root/'summary.pdf')
            with fitz.open(output) as doc:
                self.assertEqual(len(doc),4)
                expected=sorted_rows([r for r,_ in entries])
                self.assertEqual([t[1] for t in doc.get_toc()],[r.txt2 for r in expected])
                for page,row in zip(doc,expected):
                    self.assertEqual(page.get_pixmap().samples,pixels[row.txt2])
                    self.assertTrue(page.get_drawings());self.assertFalse(page.get_images())
            for path,data in before.items():self.assertEqual(path.read_bytes(),data)
            with self.assertRaises(FileExistsError):export_summary(entries,output)
            # Selecting a subset produces only those pages, not all sibling PDFs.
            with fitz.open(export_summary(entries[:1],root/'selected.pdf')) as doc:self.assertEqual(len(doc),1)
            with self.assertRaises(ValueError):export_summary([],root/'empty.pdf')

class ExcelOrderSummaryTest(unittest.TestCase):
    def test_workbook_row_order_across_models_and_selection(self):
        from openpyxl import Workbook
        from label_parser.composition import export_excel_summary
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);entries=[];pixels={}
            # Deliberately unlike alphabetical model/color or ascending size order.
            rows=[ProductionRow('',size,model,model+'-'+color+'-'+size,str(5208133300829+i),1)
                  for i,(model,color,size) in enumerate([('FBM015-099-03','LT BLUE','40'),('FBL015-166-03','BLACK','28'),('FBM015-099-03','DK BLUE','28')])]
            wb=Workbook();sheet=wb.active;sheet.append(['Article','Size','TXT1','TXT2','Barcode','Quantity'])
            for i,row in enumerate(rows):
                sheet.append([row.article,row.size,row.txt1,row.txt2,row.barcode,row.quantity]);sheet.append([None]*6)
                path=root/f'{i}.pdf'
                with fitz.open() as doc:
                    page=doc.new_page();page.insert_text((40,40),row.txt2);page.draw_rect(fitz.Rect(40,80,120+i,180))
                    doc.save(path);pixels[row.txt2]=page.get_pixmap().samples
                entries.append((row,path))
            excel=root/'manifest.xlsx';wb.save(excel);wb.close()
            output=export_excel_summary(excel,list(reversed(entries)),root/'排序汇总'/'Excel顺序组合汇总.pdf')
            with fitz.open(output) as doc:
                self.assertEqual([x[1] for x in doc.get_toc()],[r.txt2 for r in rows])
                for page,row in zip(doc,rows):
                    self.assertEqual(page.get_pixmap().samples,pixels[row.txt2]);self.assertTrue(page.get_drawings())
            with fitz.open(export_excel_summary(excel,[entries[2],entries[0]],root/'subset.pdf')) as doc:
                self.assertEqual([x[1] for x in doc.get_toc()],[rows[0].txt2,rows[2].txt2])
            with self.assertRaises(FileExistsError):export_excel_summary(excel,entries,output)
            wrong=ProductionRow('', '28','WRONG','WRONG-28','1111111111111',1)
            with self.assertRaises(ValueError):export_excel_summary(excel,[(wrong,entries[0][1])],root/'wrong.pdf')

    def test_portable_workbook_round_trip(self):
        from label_parser.composition import write_excel,read_excel
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            rows=[ProductionRow('CARELABEL_2-','28','FBM015-099-03','FBM015-099-03-DK BLUE-28','5208133300829',76)]
            path=write_excel(rows,'生产单.pdf',Path(tmp)/'生产明细.xlsx')
            self.assertEqual(read_excel(path),rows)
            book=load_workbook(path);self.assertEqual(book.active['E2'].data_type,'s');self.assertEqual(book.active.freeze_panes,'D2');book.close()
