"""Excel-driven per-Barcode compositions; preserve original PDF graphics."""
from dataclasses import dataclass
from pathlib import Path
from collections import defaultdict
import fitz
import re,zipfile
import xml.etree.ElementTree as ET
from .prodsheet import ProductionRow
from .workflow import role

SLOTS={'MAIN':(71.76,106.44),'COMPO2':(71.76,305.64),'COMPO1':(71.76,504.85),'COMPO3':(351.40,105.31),'CARE':(351.40,304.51)}

@dataclass
class Composition:
    row: ProductionRow
    matches: list
    missing: list[str]


def sorted_rows(rows):
    sizes={'XXS':0,'XS':1,'S':2,'M':3,'L':4,'XL':5,'XXL':6,'XXXL':7}
    return sorted(rows,key=lambda r:(r.txt1,r.txt2.rsplit('-',1)[0],100+int(r.size) if r.size.isdigit() else sizes.get(r.size,99),r.barcode))


def make_compositions(job,rows=None):
    rows=sorted_rows(rows if rows is not None else job.batch.details)
    index=defaultdict(list)
    for m in job.result.matches if job.result else []:
        if m.selected:index[(m.model,m.barcode,m.source)].append(m)
    items=[]
    sources=[p for p in job.sources if p in job.files_selected]
    for row in rows:
        matches=[];missing=[]
        if not sources:missing.append('未选择标签来源文件')
        missing_roles=set(SLOTS)-{role(p) for p in sources}
        if missing_roles:missing.append('缺少参考版式所需文件：'+', '.join(sorted(missing_roles)))
        used=set()
        for path in sources:
            found=index[(row.txt1,row.barcode,path)]
            kind=role(path)
            if kind not in SLOTS:missing.append(f'{Path(path).name} 没有对应版面位置');continue
            if kind in used:missing.append(f'{kind} 存在多个来源文件，无法确定唯一版面');continue
            used.add(kind)
            if len(found)!=1:missing.append(f'{Path(path).name}：匹配 {len(found)} 条，应为 1 条');continue
            matches.extend(found)
        items.append(Composition(row,matches,missing))
    return items


def compose(item,documents=None):
    if item.missing:raise ValueError('；'.join(item.missing))
    out=fitz.open();p=out.new_page(width=612,height=792)
    owned=documents is None;documents={} if owned else documents
    try:
        for m in item.matches:
            if m.source not in documents:documents[m.source]=fitz.open(m.source)
            src=documents[m.source];clip=fitz.Rect(m.rect);x,y=SLOTS[role(m.source)]
            # All original objects remain untouched; only the source viewport and placement change.
            x-=1;y-=(clip.width-2)*.065
            target=fitz.Rect(x,y,x+clip.width,y+clip.height)
            if not p.rect.contains(target):raise ValueError('标签尺寸超出参考版式，无法原尺寸组合')
            p.show_pdf_page(target,src,m.page-1,clip=clip,keep_proportion=True)
        out.set_metadata({'title':item.row.txt2,'subject':item.row.barcode,'creator':'Label Studio 0.4 — source PDF composition'})
        return out
    except Exception:
        out.close();raise
    finally:
        if owned:
            for d in documents.values():d.close()


def write_excel(rows,source,output):
    """Write a portable workbook with bundled Python dependencies only."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    output=Path(output).resolve()
    if output.exists():raise FileExistsError(f'文件已存在：{output}')
    output.parent.mkdir(parents=True,exist_ok=True)
    wb=Workbook();sheet=wb.active;sheet.title='生产明细'
    sheet.append(['Article','Size','TXT1','TXT2','Barcode','Quantity'])
    last=''
    for r in sorted_rows(rows):
        group=r.txt1+'|'+r.txt2.rsplit('-',1)[0]
        if last and group!=last:sheet.append([None]*6)
        # Explicit text type also prevents source cells being interpreted as formulas.
        sheet.append([r.article,str(r.size),r.txt1,r.txt2,str(r.barcode),r.quantity])
        for cell in list(sheet[sheet.max_row])[:5]:cell.data_type='s'
        last=group
    thin=Side(style='thin',color='DEE4EB')
    for row in sheet.iter_rows(min_col=1,max_col=6):
        sheet.row_dimensions[row[0].row].height=21
        for cell in row:
            cell.font=Font(name='Arial',size=11)
            cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
            cell.alignment=Alignment(vertical='center')
        row[1].number_format='@';row[4].number_format='@';row[5].number_format='0'
    for cell in sheet[1][:6]:
        cell.fill=PatternFill('solid',fgColor='416FBE');cell.font=Font(name='Arial',size=11,bold=True,color='FFFFFF')
    sheet.row_dimensions[1].height=27
    for col,width in [('A',20),('B',9),('C',24),('D',48),('E',24),('F',13)]:sheet.column_dimensions[col].width=width
    sheet['H1']='来源：'+Path(source).name;sheet['H1'].data_type='s'
    sheet['H1'].font=Font(name='Arial',size=10,color='64748B');sheet.column_dimensions['H'].width=65
    sheet.freeze_panes='D2';sheet.sheet_view.showGridLines=False
    wb.save(output);wb.close()
    return output


def read_excel(path):
    """Read the exported six-column manifest, not the original in-memory ordering."""
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    rows=[]
    with zipfile.ZipFile(path) as z:
        shared=[]
        if 'xl/sharedStrings.xml' in z.namelist():
            shared=[''.join(n.itertext()) for n in ET.fromstring(z.read('xl/sharedStrings.xml')).findall('s:si',ns)]
        sheet=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        for row in sheet.findall('s:sheetData/s:row',ns):
            if row.attrib['r']=='1':continue
            values={}
            for c in row.findall('s:c',ns):
                col=re.sub('[0-9]','',c.attrib['r']);v=c.find('s:v',ns)
                if c.attrib.get('t')=='inlineStr':value=''.join(c.find('s:is',ns).itertext())
                elif v is None:value=''
                elif c.attrib.get('t')=='s':value=shared[int(v.text)]
                else:value=v.text or ''
                values[col]=value
            if not values.get('C'):continue
            if not re.fullmatch(r'\d{13}',values.get('E','')):raise ValueError('Excel 中存在无效 Barcode')
            rows.append(ProductionRow(values.get('A',''),values.get('B',''),values['C'],values.get('D',''),values['E'],int(float(values.get('F','0')))))
    return rows


def safe_name(value):
    name=re.sub(r'[\\/:*?"<>|\x00-\x1f]','_',value).strip().strip('.')
    if not name:raise ValueError('TXT1/TXT2 为空，不能命名')
    return name


def export_summary(entries, output):
    """Merge explicit (ProductionRow, PDF path) pairs in color/size order."""
    entries=list(entries);output=Path(output)
    if not entries:raise ValueError('没有可汇总的组合')
    if len({r.txt1 for r,_ in entries})!=1:raise ValueError('汇总文件只能包含同一个 TXT1')
    if output.exists():raise FileExistsError(f'文件已存在：{output}')
    ordered=sorted_rows([r for r,_ in entries])
    paths={id(r):Path(p) for r,p in entries}
    with fitz.open() as merged:
        toc=[]
        for row in ordered:
            with fitz.open(paths[id(row)]) as source:
                if len(source)!=1:raise ValueError(f'{row.txt2} 必须是单页组合 PDF')
                merged.insert_pdf(source)
            toc.append([1,row.txt2,len(merged)])
        merged.set_toc(toc)
        merged.set_metadata({'title':ordered[0].txt1+'汇总PDF','creator':'LabelDesk — original PDF page merge'})
        output.parent.mkdir(parents=True,exist_ok=True)
        merged.save(output,garbage=4,deflate=True)
    return output


def export_compositions(job,folder,barcodes,progress=print):
    folder=Path(folder)
    if folder.exists() and any(folder.iterdir()):raise FileExistsError('输出文件夹不是空目录，请选择新的目录')
    folder.mkdir(parents=True,exist_ok=True)
    # The sorted workbook is the authoritative intermediate for folders/names.
    excel=write_excel(job.batch.details,job.batch.source,folder/'ProdSheet_排序明细.xlsx')
    records=read_excel(excel)
    selected=[r for r in records if r.barcode in barcodes]
    if not selected:raise ValueError('请先勾选需要导出的明细')
    items=make_compositions(job,selected)
    errors=[f'{i.row.txt2}：'+ '；'.join(i.missing) for i in items if i.missing]
    if errors:raise ValueError('已保留 Excel，但未导出不完整组合：\n'+'\n'.join(errors[:10]))
    targets=[folder/safe_name(i.row.txt1)/(safe_name(i.row.txt2)+'.pdf') for i in items]
    if len(set(targets))!=len(targets):raise ValueError('TXT2 文件名重复，未生成 PDF，请核对 Excel')
    groups=defaultdict(list)
    for item,target in zip(items,targets):groups[item.row.txt1].append((item.row,target))
    summaries={model:folder/safe_name(model)/(safe_name(model)+'汇总PDF.pdf') for model in groups}
    all_paths=targets+list(summaries.values())
    if len(set(all_paths))!=len(all_paths):raise ValueError('汇总文件名与组合文件名冲突')
    docs={}
    try:
        for index,(item,target) in enumerate(zip(items,targets)):
            target.parent.mkdir(parents=True,exist_ok=True)
            with compose(item,docs) as document:document.save(target,garbage=4,deflate=True)
            progress(f'导出组合 {index+1}/{len(items)} · {item.row.txt2}')
    finally:
        for d in docs.values():d.close()
    for model,entries in groups.items():
        export_summary(entries,summaries[model])
        progress(f'已生成 {model}汇总PDF · {len(entries)} 页')
    return excel,targets
